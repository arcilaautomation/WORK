"""Surgically repair Open_WOs.xlsm by editing its OOXML parts.

Usage: python3 build_fix.py OUTPUT.xlsm [--seed]
--seed fills the Sheet1 tab with the 216 rows from the hidden Calculation tab so the
counting formulas can be verified against known numbers (test build only).
"""
import re, sys, zipfile, datetime
from xml.sax.saxutils import escape

SRC = 'Open_WOs_original.xlsm'
OUT = sys.argv[1]
SEED = '--seed' in sys.argv

zin = zipfile.ZipFile(SRC)
parts = {n: zin.read(n) for n in zin.namelist()}
order = zin.namelist()

def get(name): return parts[name].decode('utf-8')
def put(name, text): parts[name] = text.encode('utf-8')
def sub1(text, pattern, repl, flags=0):
    new, n = re.subn(pattern, repl, text, count=1, flags=flags)
    assert n == 1, f'pattern not found: {pattern[:80]}'
    return new

def istr(ref, s, text):  # inline string cell
    return f'<c r="{ref}" s="{s}" t="inlineStr"><is><t>{escape(text)}</t></is></c>'
def num(ref, s, v): return f'<c r="{ref}" s="{s}"><v>{v}</v></c>'
def fnum(ref, s, f, v): return f'<c r="{ref}" s="{s}"><f>{escape(f)}</f><v>{v}</v></c>'
def fstr(ref, s, f, v): return f'<c r="{ref}" s="{s}" t="str"><f>{escape(f)}</f><v>{escape(v)}</v></c>'
def blank(ref, s): return f'<c r="{ref}" s="{s}"/>'
def sstr(ref, s, idx): return f'<c r="{ref}" s="{s}" t="s"><v>{idx}</v></c>'

# ---------------------------------------------------------------- workbook.xml
wb = get('xl/workbook.xml')
wb = sub1(wb, r'(<sheet name="Weekly History" sheetId="17" r:id="rId2"/>)',
          r'\1<sheet name="Sheet1" sheetId="18" r:id="rId12"/>')
wb = sub1(wb, r'<sheet name="Directions" sheetId="16" state="hidden" r:id="rId4"/>',
          '<sheet name="Directions" sheetId="16" r:id="rId4"/>')
wb = sub1(wb, r'<definedName name="_xlnm._FilterDatabase" localSheetId="2"',
          '<definedName name="_xlnm._FilterDatabase" localSheetId="3"')
wb = sub1(wb, r'<definedName name="_xleta.SUM" hidden="1" xlm="1">#NAME\?</definedName>', '')
wb = sub1(wb, r'<calcPr calcId="191029"/>', '<calcPr calcId="191029" fullCalcOnLoad="1"/>')
put('xl/workbook.xml', wb)

rels = get('xl/_rels/workbook.xml.rels')
rels = sub1(rels, r'<Relationship Id="rId9" Type="[^"]+" Target="calcChain.xml"/>', '')
rels = sub1(rels, r'</Relationships>',
            '<Relationship Id="rId12" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet5.xml"/></Relationships>')
put('xl/_rels/workbook.xml.rels', rels)

ct = get('[Content_Types].xml')
ct = sub1(ct, r'<Override PartName="/xl/calcChain.xml" ContentType="[^"]+"/>', '')
ct = sub1(ct, r'</Types>',
          '<Override PartName="/xl/worksheets/sheet5.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
          '<Override PartName="/xl/drawings/drawing2.xml" ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/></Types>')
put('[Content_Types].xml', ct)
del parts['xl/calcChain.xml']; order.remove('xl/calcChain.xml')

# ---------------------------------------------------------------- Totals (sheet1.xml)
F_BV_PM = 'COUNTIFS(Sheet1!$B:$B,"BURNSVILLE*",Sheet1!$A:$A,"PM*")'
F_BV_RT = 'COUNTIFS(Sheet1!$B:$B,"BURNSVILLE*")-' + F_BV_PM
F_LV_PM = 'COUNTIFS(Sheet1!$B:$B,"LAKEVILLE*",Sheet1!$A:$A,"PM*")'
F_LV_RT = 'COUNTIFS(Sheet1!$B:$B,"LAKEVILLE*")-' + F_LV_PM
F_STATUS = ('IF(A16=0,"Sheet1 is empty. Paste this week\'s TabWare export (All Open Work Orders & PMs) onto the Sheet1 tab and the counts fill in here.",'
            'IF(OR(Sheet1!$A$1<>"Work Order #",Sheet1!$B$1<>"Area Description"),"Check Sheet1: column A should be Work Order # and column B Area Description, with the header row in row 1.",'
            'IF(A17<>0,"Some rows on Sheet1 are not counted: their Area Description is not Burnsville or Lakeville (or they are footer rows). Check them.",'
            '"Counts are up to date with the data on Sheet1.")))')
if SEED:
    v = dict(bvrt=22, bvpm=67, bv=89, lvrt=18, lvpm=109, lv=127, all=216, n=216, nc=0,
             status='Counts are up to date with the data on Sheet1.', date='<v>46286</v>')
else:
    v = dict(bvrt=0, bvpm=0, bv=0, lvrt=0, lvpm=0, lv=0, all=0, n=0, nc=0,
             status="Sheet1 is empty. Paste this week's TabWare export (All Open Work Orders & PMs) onto the Sheet1 tab and the counts fill in here.", date='')

rows = []
rows.append(f'<row r="1">{sstr("A1",7,0)}<c r="B1" s="8">{v["date"]}</c></row>')
rows.append(f'<row r="2" ht="18.75">{blank("A2",16)}{blank("B2",17)}</row>')
rows.append(f'<row r="3" ht="15" customHeight="1">{sstr("A3",42,1)}{blank("B3",42)}</row>')
rows.append(f'<row r="4" ht="18.75">{fnum("A4",5,F_BV_RT,v["bvrt"])}{sstr("B4",6,2)}</row>')
rows.append(f'<row r="5" ht="18.75">{fnum("A5",5,F_BV_PM,v["bvpm"])}{sstr("B5",6,3)}</row>')
rows.append(f'<row r="6" ht="18.75">{fnum("A6",5,"A4+A5",v["bv"])}{sstr("B6",6,4)}</row>')
rows.append(f'<row r="7">{blank("A7",4)}</row>')
rows.append(f'<row r="8" ht="15" customHeight="1">{sstr("A8",43,5)}{blank("B8",43)}</row>')
rows.append(f'<row r="9" ht="18.75">{fnum("A9",5,F_LV_RT,v["lvrt"])}{sstr("B9",6,2)}</row>')
rows.append(f'<row r="10" ht="18.75">{fnum("A10",5,F_LV_PM,v["lvpm"])}{sstr("B10",6,3)}</row>')
rows.append(f'<row r="11" ht="18.75">{fnum("A11",5,"A9+A10",v["lv"])}{sstr("B11",6,4)}</row>')
rows.append(f'<row r="13" ht="18.75">{sstr("A13",44,6)}{blank("B13",44)}</row>')
rows.append(f'<row r="14" ht="18.75">{fnum("A14",5,"A6+A11",v["all"])}{sstr("B14",6,7)}</row>')
rows.append(f'<row r="16">{fnum("A16",3,"MAX(0,COUNTA(Sheet1!$A:$A)-1)",v["n"])}{istr("B16",4,"WORK ORDERS ON SHEET1 (ALL AREAS)")}</row>')
rows.append(f'<row r="17">{fnum("A17",3,"A16-A14",v["nc"])}{istr("B17",4,"NOT COUNTED (AREA IS NOT BURNSVILLE OR LAKEVILLE)")}</row>')
rows.append(f'<row r="19" ht="30" customHeight="1">{fstr("A19",11,F_STATUS,v["status"])}</row>')
sheetdata = '<sheetData>' + ''.join(rows) + '</sheetData>'

s1 = get('xl/worksheets/sheet1.xml')
s1 = sub1(s1, r'<dimension ref="A1:N21"/>', '<dimension ref="A1:E19"/>')
s1 = sub1(s1, r'<selection activeCell="H19" sqref="H19"/>', '<selection activeCell="A1" sqref="A1"/>')
s1 = sub1(s1, r'<sheetData>.*?</sheetData>', sheetdata, flags=re.S)
s1 = sub1(s1, r'<mergeCells count="9">.*?</mergeCells>',
          '<mergeCells count="4"><mergeCell ref="A3:B3"/><mergeCell ref="A8:B8"/><mergeCell ref="A13:B13"/><mergeCell ref="A19:E19"/></mergeCells>', flags=re.S)
cf = ('<conditionalFormatting sqref="A17:B17"><cfRule type="expression" dxfId="13" priority="1"><formula>$A$17&lt;&gt;0</formula></cfRule></conditionalFormatting>'
      '<conditionalFormatting sqref="A19"><cfRule type="expression" dxfId="11" priority="2"><formula>LEFT($A$19,6)="Counts"</formula></cfRule>'
      '<cfRule type="expression" dxfId="14" priority="3"><formula>OR(LEFT($A$19,5)="Check",LEFT($A$19,4)="Some")</formula></cfRule></conditionalFormatting>')
s1 = sub1(s1, r'<conditionalFormatting sqref="A16">.*?<conditionalFormatting sqref="M16">.*?</conditionalFormatting>', cf, flags=re.S)
s1 = sub1(s1, r'<legacyDrawing r:id="rId2"/>', '<drawing r:id="rId4"/><legacyDrawing r:id="rId2"/>')
put('xl/worksheets/sheet1.xml', s1)

r1 = get('xl/worksheets/_rels/sheet1.xml.rels')
r1 = sub1(r1, r'</Relationships>',
          '<Relationship Id="rId4" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing" Target="../drawings/drawing2.xml"/></Relationships>')
put('xl/worksheets/_rels/sheet1.xml.rels', r1)

drawing2 = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
 '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
 '<xdr:twoCellAnchor><xdr:from><xdr:col>3</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>0</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>'
 '<xdr:to><xdr:col>4</xdr:col><xdr:colOff>640080</xdr:colOff><xdr:row>2</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>'
 '<xdr:sp macro="[0]!UpdateData" textlink=""><xdr:nvSpPr><xdr:cNvPr id="2" name="Refresh Button"/><xdr:cNvSpPr/></xdr:nvSpPr>'
 '<xdr:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/></a:xfrm><a:prstGeom prst="roundRect"><a:avLst/></a:prstGeom>'
 '<a:solidFill><a:srgbClr val="2A78D6"/></a:solidFill><a:ln><a:noFill/></a:ln></xdr:spPr>'
 '<xdr:txBody><a:bodyPr vertOverflow="clip" horzOverflow="clip" wrap="square" rtlCol="0" anchor="ctr"/><a:lstStyle/>'
 '<a:p><a:pPr algn="ctr"/><a:r><a:rPr lang="en-US" sz="1100" b="1"><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill></a:rPr><a:t>Refresh</a:t></a:r></a:p></xdr:txBody>'
 '</xdr:sp><xdr:clientData/></xdr:twoCellAnchor></xdr:wsDr>')
put('xl/drawings/drawing2.xml', drawing2); order.append('xl/drawings/drawing2.xml')

# comments: keep A4 only (D4 and M4 blocks are removed)
c1 = get('xl/comments1.xml')
c1 = sub1(c1, r'<comment ref="D4".*?</comment>', '', flags=re.S)
c1 = sub1(c1, r'<comment ref="M4".*?</comment>', '', flags=re.S)
put('xl/comments1.xml', c1)
vml = get('xl/drawings/vmlDrawing1.vml')
vml = sub1(vml, r'<v:shape id="_x0000_s1026".*?</v:shape>', '', flags=re.S)
vml = sub1(vml, r'<v:shape id="_x0000_s1027".*?</v:shape>', '', flags=re.S)
put('xl/drawings/vmlDrawing1.vml', vml)

# ---------------------------------------------------------------- Weekly History (sheet2.xml) + table + chart
row23 = ('<row r="23"><c r="A23" s="39"><v>46286</v></c><c r="B23" s="38"><v>6</v></c><c r="C23" s="38"><v>6</v></c>'
 '<c r="D23" s="40"><f>WeeklyLog[[#This Row],[Burnsville Repair Tickets]]+WeeklyLog[[#This Row],[Burnsville PMs]]</f><v>12</v></c>'
 '<c r="E23" s="38"><v>38</v></c><c r="F23" s="38"><v>240</v></c>'
 '<c r="G23" s="40"><f>WeeklyLog[[#This Row],[Lakeville Repair Tickets]]+WeeklyLog[[#This Row],[Lakeville PMs]]</f><v>278</v></c>'
 '<c r="H23" s="40"><f>WeeklyLog[[#This Row],[Burnsville Total]]+WeeklyLog[[#This Row],[Lakeville Total]]</f><v>290</v></c>'
 '<c r="I23" s="41"><f ca="1">IFERROR(WeeklyLog[[#This Row],[Combined Total]]-OFFSET(WeeklyLog[[#This Row],[Combined Total]],-1,0),"")</f><v>71</v></c></row>')
s2 = get('xl/worksheets/sheet2.xml')
s2 = sub1(s2, r'<dimension ref="A1:I22"/>', '<dimension ref="A1:I23"/>')
s2 = sub1(s2, r'</row></sheetData>', '</row>' + row23 + '</sheetData>')
s2 = sub1(s2, r'<conditionalFormatting sqref="I22">', '<conditionalFormatting sqref="I22:I1000">')
put('xl/worksheets/sheet2.xml', s2)
t1 = get('xl/tables/table1.xml')
t1 = t1.replace('ref="A21:I22"', 'ref="A21:I23"'); assert t1.count('ref="A21:I23"') == 2
put('xl/tables/table1.xml', t1)

ch = get('xl/charts/chart1.xml')
sers = ch.split('<c:ser>')
assert len(sers) == 7
vals = [12, 278, 290, 12, 278, 290]
for i in range(1, 7):
    s = sers[i]
    s = sub1(s, r'<c:ptCount val="1"/><c:pt idx="0"><c:v>46280</c:v></c:pt>',
             '<c:ptCount val="2"/><c:pt idx="0"><c:v>46280</c:v></c:pt><c:pt idx="1"><c:v>46286</c:v></c:pt>')
    if i <= 3:   # visible lines keep point 0 and gain point 1
        s = sub1(s, r'(<c:formatCode>#,##0</c:formatCode>)<c:ptCount val="1"/>(<c:pt idx="0"><c:v>\d+</c:v></c:pt>)',
                 r'\1<c:ptCount val="2"/>\2' + f'<c:pt idx="1"><c:v>{vals[i-1]}</c:v></c:pt>')
    else:        # last-point label series: point 0 is now #N/A, only point 1 has a value
        s = sub1(s, r'(<c:formatCode>#,##0</c:formatCode>)<c:ptCount val="1"/><c:pt idx="0"><c:v>\d+</c:v></c:pt>',
                 r'\1<c:ptCount val="2"/>' + f'<c:pt idx="1"><c:v>{vals[i-1]}</c:v></c:pt>')
    sers[i] = s
put('xl/charts/chart1.xml', '<c:ser>'.join(sers))

# ---------------------------------------------------------------- Directions (sheet4.xml)
def box_single(r, n, text):
    return (f'<row r="{r}">{num(f"A{r}",26,n)}{istr(f"B{r}",19,text)}{blank(f"C{r}",19)}{blank(f"D{r}",19)}'
            f'{blank(f"E{r}",19)}{blank(f"F{r}",19)}{blank(f"G{r}",20)}</row>')
def box_multi(r0, n, title, lines):
    out = [f'<row r="{r0}">' + (num(f"A{r0}",45,n) if n else blank(f"A{r0}",45)) +
           f'{istr(f"B{r0}",27,title)}{blank(f"C{r0}",21)}{blank(f"D{r0}",21)}{blank(f"E{r0}",21)}{blank(f"F{r0}",21)}{blank(f"G{r0}",22)}</row>']
    r = r0
    for j, line in enumerate(lines):
        r += 1
        last = (j == len(lines) - 1)
        if last:
            out.append(f'<row r="{r}">{blank(f"A{r}",47)}{blank(f"B{r}",29)}{istr(f"C{r}",24,line)}{blank(f"D{r}",24)}{blank(f"E{r}",24)}{blank(f"F{r}",24)}{blank(f"G{r}",25)}</row>')
        else:
            out.append(f'<row r="{r}">{blank(f"A{r}",46)}{blank(f"B{r}",28)}{istr(f"C{r}",4,line)}{blank(f"G{r}",23)}</row>')
    return ''.join(out), r
def spacer(r): return f'<row r="{r}" ht="9" customHeight="1"/>'

d = []; merges = ['A1:G1']
d.append(f'<row r="1" ht="26.25">{istr("A1",48,"DIRECTIONS")}{blank("B1",49)}{blank("C1",49)}{blank("D1",49)}{blank("E1",49)}{blank("F1",49)}{blank("G1",50)}</row>')
d.append(spacer(2))
d.append(box_single(3, 1, 'Clear the Sheet1 tab: click the Sheet1 tab, press Ctrl+A, then press Delete.'))
d.append(spacer(4))
blk, r = box_multi(5, 2, "Download this week's data from TabWare:",
    ['TabWare', 'Work Order Search', 'Query: All Open Work Orders & PMs', 'Search', 'File', 'Save As',
     'Save In: Downloads', 'File type: Excel 2007', 'Save']); d.append(blk); merges.append(f'A5:A{r}')
d.append(spacer(r+1)); r += 2
d.append(box_single(r, 3, 'Open the downloaded file.'))
d.append(spacer(r+1)); r += 2
blk, r2 = box_multi(r, 4, 'Put the new data on the Sheet1 tab of this file:',
    ['In the downloaded file press Ctrl+A, then Ctrl+C.',
     'On the Sheet1 tab click cell A1 and paste as values only',
     '(right-click > Paste Values). Keep the header row in row 1.']); d.append(blk); merges.append(f'A{r}:A{r2}'); r = r2
d.append(spacer(r+1)); r += 2
blk, r2 = box_multi(r, 5, 'Check the Totals tab:',
    ['The counts update on their own from Sheet1. There is nothing to run.',
     'Click the Refresh button to stamp today\'s date in B1, or type the date.',
     'The NOT COUNTED line should read 0. If it does not, check the',
     'Area Description column on Sheet1: only Burnsville and Lakeville count.']); d.append(blk); merges.append(f'A{r}:A{r2}'); r = r2
d.append(spacer(r+1)); r += 2
blk, r2 = box_multi(r, 6, 'Add this week to the Weekly History tab:',
    ['Type the date and the four counts from Totals (repair tickets and',
     'PMs for each plant) in the next empty row of the table.',
     'The totals, the change vs last week, and the chart fill in on their own.']); d.append(blk); merges.append(f'A{r}:A{r2}'); r = r2
d.append(spacer(r+1)); r += 2
d.append(box_single(r, 7, 'Save the file.'))
d.append(spacer(r+1)); r += 2
blk, r2 = box_multi(r, None, 'Notes:',
    ['Refresh runs a macro. If Excel blocks it, right-click the file in File Explorer,',
     'choose Properties, tick Unblock, and reopen it. The counts work without macros.',
     'The Refresh macro also copies Sheet1 to the hidden Calculation tab.']); d.append(blk); merges.append(f'A{r}:A{r2}'); r = r2
s4 = get('xl/worksheets/sheet4.xml')
s4 = sub1(s4, r'<dimension ref="A1:G29"/>', f'<dimension ref="A1:G{r}"/>')
s4 = sub1(s4, r'<cols>.*?</cols>', '<cols><col min="1" max="1" width="9.140625" style="3"/><col min="2" max="7" width="13.7109375" style="4" customWidth="1"/><col min="8" max="16384" width="9.140625" style="4"/></cols>', flags=re.S)
s4 = sub1(s4, r'<sheetData>.*?</sheetData>', '<sheetData>' + ''.join(d) + '</sheetData>', flags=re.S)
s4 = sub1(s4, r'<mergeCells count="4">.*?</mergeCells>',
          f'<mergeCells count="{len(merges)}">' + ''.join(f'<mergeCell ref="{m}"/>' for m in merges) + '</mergeCells>', flags=re.S)
put('xl/worksheets/sheet4.xml', s4)

# ---------------------------------------------------------------- Sheet1 (new sheet5.xml)
headers = ['Work Order #','Area Description','Priority','Entered Date','Requested Completion','Status Description',
           'Work Order Description','Equipment Description','Long Description','Requestor','Assigned To Name',
           'Craft Actual Hours','Completion Remarks']
widths = [9.85546875,13.42578125,9.85546875,10.28515625,13.7109375,13.42578125,25.140625,25.28515625,29.42578125,12.42578125,11.28515625,8.85546875,64.28515625]
cols = ''.join(f'<col min="{i+1}" max="{i+1}" width="{w}" customWidth="1"/>' for i, w in enumerate(widths))
letters = 'ABCDEFGHIJKLM'
hdr = ''.join(istr(f'{letters[i]}1', 2, h) for i, h in enumerate(headers))
srows = [f'<row r="1" ht="45" customHeight="1">{hdr}</row>']
maxrow = 1
if SEED:
    from openpyxl import load_workbook
    wsC = load_workbook(SRC, keep_vba=True, data_only=True)['Calculation']
    r = 1
    for rec in wsC.iter_rows(min_row=3, max_row=wsC.max_row, values_only=True):
        if rec[1] is None: continue
        r += 1
        cells = []
        for i, val in enumerate(rec[1:14]):
            ref = f'{letters[i]}{r}'
            if val is None: continue
            if isinstance(val, datetime.datetime):
                cells.append(num(ref, 0, (val - datetime.datetime(1899,12,30)).days))
            elif isinstance(val, (int, float)):
                cells.append(num(ref, 0, val))
            else:
                cells.append(istr(ref, 0, str(val)))
        srows.append(f'<row r="{r}">' + ''.join(cells) + '</row>')
    maxrow = r
sheet5 = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
 '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
 '<sheetPr><tabColor rgb="FFFFC000"/></sheetPr>'
 f'<dimension ref="A1:M{maxrow}"/>'
 '<sheetViews><sheetView workbookViewId="0"><pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen"/>'
 '<selection pane="bottomLeft" activeCell="A2" sqref="A2"/></sheetView></sheetViews>'
 '<sheetFormatPr defaultRowHeight="15"/>'
 f'<cols>{cols}</cols>'
 '<sheetData>' + ''.join(srows) + '</sheetData>'
 '<pageMargins left="0.7" right="0.7" top="0.75" bottom="0.75" header="0.3" footer="0.3"/></worksheet>')
put('xl/worksheets/sheet5.xml', sheet5); order.append('xl/worksheets/sheet5.xml')

# ---------------------------------------------------------------- docProps/app.xml
app = get('docProps/app.xml')
app = sub1(app, r'<vt:lpstr>Worksheets</vt:lpstr></vt:variant><vt:variant><vt:i4>4</vt:i4>',
           '<vt:lpstr>Worksheets</vt:lpstr></vt:variant><vt:variant><vt:i4>5</vt:i4>')
app = sub1(app, r'<vt:vector size="8" baseType="lpstr"><vt:lpstr>Totals</vt:lpstr><vt:lpstr>Weekly History</vt:lpstr>',
           '<vt:vector size="9" baseType="lpstr"><vt:lpstr>Totals</vt:lpstr><vt:lpstr>Weekly History</vt:lpstr><vt:lpstr>Sheet1</vt:lpstr>')
put('docProps/app.xml', app)

# ---------------------------------------------------------------- write
with zipfile.ZipFile(OUT, 'w', zipfile.ZIP_DEFLATED) as zout:
    names = ['[Content_Types].xml'] + [n for n in order if n != '[Content_Types].xml']
    for n in names:
        zout.writestr(n, parts[n])
print('wrote', OUT, 'seed' if SEED else '', 'parts:', len(names))
