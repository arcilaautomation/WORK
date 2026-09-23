#!/usr/bin/env python3
"""
Build Open_WOs_weekly_log.xlsx — the weekly log on its own.

The boss's instruction, on Open_WOs_w_graph.xlsx: "We only need Lines 21, 22,
and 23 from sheet 1." Those lines are the WeeklyLog table: its header row and
the weekly rows under it. Everything else goes — the title, the graph, the
graph's named ranges and the Totals tab.

This transforms that Excel-saved workbook rather than rebuilding it, so every
colour, font, number format, the custom "Weekly Log Style" table style and the
table's calculated columns stay exactly as Excel wrote them. The table moves to
the top of the sheet; nothing about how it looks changes.

Weekly routine for whoever keeps it: type the date and the four counts in the
row under the last week. Excel grows the table and fills in the three totals
and the change on its own, because those columns are calculated columns.

Run:  python3 build_weekly_log.py [source.xlsx] [output.xlsx]
"""
import re
import shutil
import sys
import zipfile
from xml.sax.saxutils import escape

from openpyxl import load_workbook
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.utils import range_boundaries

SOURCE = "source/Open_WOs_w_graph.xlsx"
OUTPUT = "Open_WOs_weekly_log.xlsx"
TABLE = "WeeklyLog"
SHEET_NAME = "Weekly Log"

# Columns of the WeeklyLog table, as they sit in the source.
TYPED = ("B", "C", "E", "F")            # the four counts, typed each week
BV_TOT, LV_TOT, COMB, CHANGE = "D", "G", "H", "I"


def find_table(wb):
    for ws in wb.worksheets:
        if TABLE in ws.tables:
            return ws, ws.tables[TABLE]
    raise SystemExit(f"no table named {TABLE} in the source workbook")


def build(src, out):
    wb = load_workbook(src)
    ws, table = find_table(wb)
    min_col, head, max_col, last = range_boundaries(table.ref)
    shift = head - 1                                  # header lands on row 1
    head_height = ws.row_dimensions[head].height

    # ---- only the log is wanted -------------------------------------------
    for name in list(wb.sheetnames):
        if name != ws.title:
            wb.remove(wb[name])                       # the Totals tab
    ws._charts = []                                   # the graph
    for name in [n for n in ws.defined_names if n.startswith(("chart_", "last_"))]:
        del ws.defined_names[name]                    # ranges only the graph used
    for row in ws.iter_rows(min_row=1, max_row=head - 1):
        for cell in row:
            cell.value = None                         # the title above the table

    # ---- lines 21-23 move to the top, formats and all ------------------------
    first_col, last_col = ws.cell(1, min_col).column_letter, ws.cell(1, max_col).column_letter
    ws.move_range(f"{first_col}{head}:{last_col}{last}", rows=-shift, translate=True)
    new_ref = f"{first_col}1:{last_col}{last - shift}"
    table.ref = new_ref
    if table.autoFilter is not None:
        table.autoFilter.ref = new_ref
    for r in list(ws.row_dimensions):
        if r >= 1:
            ws.row_dimensions[r].height = None
    ws.row_dimensions[1].height = head_height

    # ---- the red/green change colours follow the column up -------------------
    rules = []
    for cf in ws.conditional_formatting:
        for rule in cf.rules:
            rules.append((str(cf.sqref), rule))
    ws.conditional_formatting = ConditionalFormattingList()
    for sqref, rule in rules:
        def up(m):
            return f"{m.group(1)}{int(m.group(2)) - shift}"
        first_row = int(re.match(r"[A-Z]+(\d+)", sqref).group(1))
        rule.formula = [re.sub(r"\b([A-Z]+)(\d+)\b",
                               lambda m: up(m) if int(m.group(2)) == first_row else m.group(0),
                               f) for f in rule.formula]
        new_sqref = re.sub(r"^([A-Z]+)(\d+)", lambda m: up(m), sqref)
        ws.conditional_formatting.add(new_sqref, rule)

    # ---- small conveniences that do not change how it looks ------------------
    ws.title = SHEET_NAME
    ws.freeze_panes = "A2"                            # header stays put as it grows
    next_row = last - shift + 1
    for sel in ws.sheet_view.selection:
        sel.activeCell = f"A{next_row}"               # opens on next week's row
        sel.sqref = f"A{next_row}"
    ws.print_title_rows = "1:1"                       # header repeats when printed
    wb.calculation.fullCalcOnLoad = True
    wb.active = 0
    wb.save(out)

    store_results(out, SHEET_NAME, 2, last - shift)
    print(f"wrote {out}: table {new_ref}, {last - head} weeks, sheet '{SHEET_NAME}'")


def store_results(path, sheet, first, last):
    """Put each total's value beside its formula, as Excel does when it saves.

    openpyxl keeps the formulas but drops the results Excel had stored, so a
    preview that does not calculate (Outlook, iPhone Quick Look, SharePoint
    thumbnails) would show blank totals. The results are simple sums of the
    typed counts, computed here from the same cells the formulas read.
    """
    wb = load_workbook(path)
    ws = wb[sheet]
    results, prev = {}, None
    for r in range(first, last + 1):
        b, c, e, f = (ws[f"{col}{r}"].value for col in TYPED)
        bvt, lvt = b + c, e + f
        comb = bvt + lvt
        results[f"{BV_TOT}{r}"] = bvt
        results[f"{LV_TOT}{r}"] = lvt
        results[f"{COMB}{r}"] = comb
        # first row: OFFSET lands on the header text, IFERROR returns ""
        results[f"{CHANGE}{r}"] = "" if prev is None else comb - prev
        prev = comb

    with zipfile.ZipFile(path) as z:
        parts = {i.filename: z.read(i.filename) for i in z.infolist()}
        infos = {i.filename: i for i in z.infolist()}
    part = "xl/worksheets/sheet1.xml"
    xml = parts[part].decode()

    def fill(m):
        coord, attrs, formula = m.group(1), m.group(2), m.group(3)
        if coord not in results:
            return m.group(0)
        v = results[coord]
        attrs = re.sub(r'\s+t="[^"]*"', "", attrs)
        if isinstance(v, str):
            attrs += ' t="str"'
        return f'<c r="{coord}"{attrs}><f>{formula}</f><v>{escape(str(v))}</v></c>'

    xml = re.sub(r'<c r="([A-Z]+\d+)"([^>]*)><f>(.*?)</f><v\s*/></c>', fill, xml,
                 flags=re.S)
    parts[part] = xml.encode()
    tmp = path + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as outzip:
        for name, data in parts.items():
            outzip.writestr(infos[name], data)
    shutil.move(tmp, path)
    print(f"stored results for {len(results)} formula cells")


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else SOURCE,
          sys.argv[2] if len(sys.argv) > 2 else OUTPUT)
